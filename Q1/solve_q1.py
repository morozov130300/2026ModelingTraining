"""
问题一：微网日前购电策略建模与求解
====================================
严格按《问题一_日前购电策略建模思路.md》实现：

  1. 文档第4节 基准 MILP（含充放电互斥二元变量 u_t），采用分支切割法
     （scipy.optimize.milp，底层 HiGHS）求解；
  2. 文档5.3节 LP 主模型（删除 u_t，0<=P_ch,P_dis<=5000）；
  3. 文档5.4节 两阶段词典序优化：第一阶段求最小购电成本 C*，
     第二阶段在 C <= C* + eps 下最小化储能吞吐量 Σ(P_ch+P_dis)Δt；
  4. 文档5.5节 验证：充放电互斥 min(P_ch,P_dis) <= τ=1e-6、功率平衡、
     光伏平衡、SOC 范围与日末 SOC；比较 LP 与 MILP 最优成本；
  5. 文档第6节 输出：全天购电量、购电费、弃光量、储能吞吐量、
     最大购电功率，并按附件5模板原结构填写 result1.xlsx。

模型变量（与文档4.7一致）：
    P_buy, P_ch, P_dis, E(1..145), P_pv_use, P_cur,（MILP 另含 u）
其中 P_pv_use、P_cur 作为显式变量，用等式约束建模（不用消元，
避免符号错误），非负性由变量下界保证。

用法：
    python solve_q1.py
"""

import os
import sys
import time

import numpy as np
import openpyxl
from openpyxl import load_workbook
from scipy.optimize import linprog, milp, LinearConstraint, Bounds

# ========== 路径 ==========
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)
ATTACH_DIR = os.path.join(PROJECT_ROOT, '附件')
DATA_FILE = os.path.join(ATTACH_DIR, '附件1.xlsx')
TEMPLATE_FILE = os.path.join(ATTACH_DIR, '附件5', 'result1.xlsx')
OUT_FILE = os.path.join(SCRIPT_DIR, 'result1.xlsx')
TABLE1_FILE = os.path.join(SCRIPT_DIR, '表1.xlsx')
TABLE2_FILE = os.path.join(SCRIPT_DIR, '表2.xlsx')

# ========== 已知参数（文档第2、3节） ==========
DT = 10 / 60                      # Δt = 1/6 h
ETA_CH = 0.9                      # η^ch
ETA_DIS = 0.9                     # η^dis
E_MIN, E_MAX = 1200.0, 10800.0    # 储能运行电量边界 (kWh)
E_INIT = 6000.0                   # E_1 (kWh)
E_END = 6000.0                    # E_145 (kWh)
P_CH_MAX = 5000.0                 # 最大充电功率 (kW)
P_DIS_MAX = 5000.0                # 最大放电功率 (kW)
# P^buy,max（文档3.1）：题目未提供变压器容量/馈线限额/购电合同数值，
# 理论上无上限约束。
P_BUY_MAX = np.inf

TAU = 1e-6        # 文档5.5：充放电互斥校验阈值 τ (kW)
BAL_TOL = 1e-4    # 等式约束校验容差（功率 kW / 电量 kWh）


# ========== 1. 读取附件1数据（文档第2节：144个10分钟时段） ==========
def read_data(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    prices, loads, pv = [], [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] is None:
            continue
        prices.append(float(row[1]))
        loads.append(float(row[2]))
        pv.append(float(row[3]))
    wb.close()
    if len(prices) != 144:
        raise ValueError(f"期望144个时段，实际{len(prices)}个")

    N = len(prices)
    prices_arr = np.asarray(prices)
    loads_arr = np.asarray(loads)
    pv_arr = np.asarray(pv)

    # 数据基本统计
    print(f"\n[数据读取] {len(prices)} 个时段（Δt={DT:.4f} h）")
    print(f"  电价: [{prices_arr.min():.4f}, {prices_arr.max():.4f}] 元/kWh, 均值={prices_arr.mean():.4f}")
    print(f"  负荷: [{loads_arr.min():.1f}, {loads_arr.max():.1f}] kW, 均值={loads_arr.mean():.1f}")
    print(f"  光伏: [{pv_arr.min():.1f}, {pv_arr.max():.1f}] kW, 均值={pv_arr.mean():.1f}")

    # 能量统计
    total_load = float(np.sum(loads_arr) * DT)
    total_pv = float(np.sum(pv_arr) * DT)
    net_load = loads_arr - pv_arr
    total_net = float(np.sum(np.maximum(net_load, 0)) * DT)  # 净购电需求

    surplus_periods = [(t, pv_arr[t] - loads_arr[t]) for t in range(N) if pv_arr[t] > loads_arr[t]]
    deficit_periods = [(t, loads_arr[t] - pv_arr[t]) for t in range(N) if pv_arr[t] < loads_arr[t]]

    print(f"\n[数据统计]")
    print(f"  全天负载电量: {total_load/1000:.3f} MWh")
    print(f"  全天预测光伏: {total_pv/1000:.3f} MWh")
    print(f"  净购电需求: {total_net:.2f} kWh")
    print(f"  光伏盈余时段: {len(surplus_periods)}个, 总盈余={sum(s[1] for s in surplus_periods)*DT:.2f} kWh")
    print(f"  光伏 deficit 时段: {len(deficit_periods)}个, 总 deficit={sum(d[1] for d in deficit_periods)*DT:.2f} kWh")
    print(f"  最大净负荷: {np.max(net_load):.1f} kW (时段{np.argmax(net_load)})")

    return prices, loads, pv


# ========== 2. 约束构建（文档4.7完整模型，MILP/LP共用变量布局） ==========
# 变量 x = [P_buy(0..N-1), P_ch, P_dis, E(0..N), P_pv_use, P_cur, (u 仅MILP)]
# 约定：E[t] 对应文档记号 E_{t+1}（t=0..N，共145个状态点）
def build_constraints(N, loads, pv, with_u):
    i_buy = lambda t: t
    i_ch = lambda t: N + t
    i_dis = lambda t: 2 * N + t
    i_E = lambda t: 3 * N + t
    i_pvu = lambda t: 4 * N + 1 + t
    i_cur = lambda t: 5 * N + 1 + t
    i_u = lambda t: 6 * N + 1 + t
    n = (7 * N + 1) if with_u else (6 * N + 1)

    A_eq, b_eq, A_ub, b_ub = [], [], [], []

    for t in range(N):
        # 文档4.2 功率平衡：P_buy + P_pv_use + P_dis = L + P_ch
        r = np.zeros(n)
        r[i_buy(t)] = 1.0
        r[i_pvu(t)] = 1.0
        r[i_dis(t)] = 1.0
        r[i_ch(t)] = -1.0
        A_eq.append(r); b_eq.append(float(loads[t]))

        # 文档4.3 光伏平衡：P_pv_use + P_cur = P_pv_hat
        r = np.zeros(n)
        r[i_pvu(t)] = 1.0
        r[i_cur(t)] = 1.0
        A_eq.append(r); b_eq.append(float(pv[t]))

        # 文档4.4 储能状态转移：
        # E_{t+1} = E_t + η_ch·P_ch·Δt − P_dis·Δt/η_dis
        r = np.zeros(n)
        r[i_E(t + 1)] = 1.0
        r[i_E(t)] = -1.0
        r[i_ch(t)] = -ETA_CH * DT
        r[i_dis(t)] = DT / ETA_DIS
        A_eq.append(r); b_eq.append(0.0)

    # 文档4.4 日初、日末SOC：E_1 = E_145 = 6000
    r = np.zeros(n); r[i_E(0)] = 1.0
    A_eq.append(r); b_eq.append(E_INIT)
    r = np.zeros(n); r[i_E(N)] = 1.0
    A_eq.append(r); b_eq.append(E_END)

    if with_u:
        # 文档4.5 充放电互斥：P_ch <= 5000·u_t，P_dis <= 5000·(1−u_t)
        for t in range(N):
            r = np.zeros(n)
            r[i_ch(t)] = 1.0
            r[i_u(t)] = -P_CH_MAX
            A_ub.append(r); b_ub.append(0.0)
            r = np.zeros(n)
            r[i_dis(t)] = 1.0
            r[i_u(t)] = P_DIS_MAX
            A_ub.append(r); b_ub.append(P_DIS_MAX)

    return {
        'n': n, 'N': N,
        'A_eq': np.array(A_eq), 'b_eq': np.array(b_eq),
        'A_ub': np.array(A_ub) if A_ub else None,
        'b_ub': np.array(b_ub) if A_ub else None,
        'i_buy': i_buy, 'i_ch': i_ch, 'i_dis': i_dis, 'i_E': i_E,
        'i_pvu': i_pvu, 'i_cur': i_cur, 'i_u': i_u,
    }


def build_bounds(model, with_u):
    """变量边界：文档4.3(P_pv_use,P_cur>=0)、4.4(E边界)、4.6(P_buy<=P_buy_max)、5.3(LP功率上界)"""
    N = model['N']
    lo = np.zeros(model['n'])
    hi = np.full(model['n'], np.inf)
    hi[:N] = P_BUY_MAX                     # 文档4.6 购电功率上界
    hi[N:2 * N] = P_CH_MAX                 # 文档5.3（LP形式）
    hi[2 * N:3 * N] = P_DIS_MAX
    lo[3 * N:4 * N + 1] = E_MIN            # 文档4.4
    hi[3 * N:4 * N + 1] = E_MAX
    # P_pv_use、P_cur 下界0、上界不设（上界由光伏平衡等式与非负弃光隐含）
    if with_u:
        lo[6 * N + 1:] = 0.0               # 文档4.5 u_t ∈ {0,1}
        hi[6 * N + 1:] = 1.0
    return lo, hi


def cost_vector(model, prices):
    """文档4.1 目标函数：min Σ c_t·P_buy·Δt"""
    c = np.zeros(model['n'])
    c[:model['N']] = np.asarray(prices) * DT
    return c


def throughput_vector(model):
    """文档5.4 第二阶段目标：min Σ (P_ch + P_dis)·Δt"""
    c = np.zeros(model['n'])
    N = model['N']
    c[N:3 * N] = DT
    return c


def extract(x, model):
    N = model['N']
    return {
        'P_buy': x[:N],
        'P_ch': x[N:2 * N],
        'P_dis': x[2 * N:3 * N],
        'E': x[3 * N:4 * N + 1],
        'P_pvu': x[4 * N + 1:5 * N + 1],
        'P_cur': x[5 * N + 1:6 * N + 1],
    }


# ========== 3. 求解（文档5.6 步骤3-5） ==========
def solve_milp(model, prices):
    """文档第4节基准MILP：分支切割法（scipy.optimize.milp → HiGHS）"""
    lo, hi = build_bounds(model, with_u=True)
    integrality = np.zeros(model['n'])
    integrality[6 * model['N'] + 1:] = 1
    cons = [LinearConstraint(model['A_eq'], model['b_eq'], model['b_eq'])]
    cons.append(LinearConstraint(model['A_ub'], -np.inf, model['b_ub']))
    t0 = time.perf_counter()
    res = milp(c=cost_vector(model, prices), constraints=cons,
               integrality=integrality, bounds=Bounds(lo, hi),
               options={'mip_rel_gap': 1e-9, 'time_limit': 600})
    elapsed = time.perf_counter() - t0
    if not res.success:
        raise RuntimeError(f"MILP求解失败: {res.message}")
    print(f"\n[MILP求解]")
    print(f"  最优费用: {res.fun:.4f} 元")
    print(f"  求解耗时: {elapsed:.2f} s")
    print(f"  变量总数: {model['n']} (含{model['N']}个二元变量)")
    return res, elapsed


def solve_lp_two_stage(model, prices):
    """文档5.3 LP主模型 + 文档5.4 两阶段词典序优化"""
    lo, hi = build_bounds(model, with_u=False)
    bounds = list(zip(lo, hi))
    c_cost = cost_vector(model, prices)

    # 第一阶段：C* = min C
    t0 = time.perf_counter()
    res1 = linprog(c_cost, A_eq=model['A_eq'], b_eq=model['b_eq'],
                   bounds=bounds, method='highs')
    t1 = time.perf_counter() - t0
    if not res1.success:
        raise RuntimeError(f"LP第一阶段求解失败: {res1.message}")
    c_star = float(res1.fun)
    print(f"\n[LP第一阶段]")
    print(f"  最优费用 C* = {c_star:.4f} 元")
    print(f"  求解耗时: {t1:.2f} s")

    # 第二阶段：C <= C* + ε 下 min Σ(P_ch+P_dis)Δt
    # ε 取与求解器可行性容差一致的极小正数（取 C* 的 1e-6 相对量级）
    eps = 1e-6 * max(1.0, abs(c_star))
    t0 = time.perf_counter()
    res2 = linprog(throughput_vector(model),
                   A_ub=c_cost.reshape(1, -1), b_ub=np.array([c_star + eps]),
                   A_eq=model['A_eq'], b_eq=model['b_eq'],
                   bounds=bounds, method='highs')
    t2 = time.perf_counter() - t0
    if not res2.success:
        raise RuntimeError(f"LP第二阶段求解失败: {res2.message}")
    actual_cost = float(c_cost @ res2.x)
    print(f"\n[LP第二阶段]")
    print(f"  储能吞吐量: {res2.fun:.4f} kWh")
    print(f"  实际费用: {actual_cost:.4f} 元 (<= C*+{eps:.2e})")
    print(f"  求解耗时: {t2:.2f} s")
    return res1, res2, c_star, eps, t1, t2


# ========== 4. 验证（文档5.5、第6节） ==========
def validate(sol, model, loads, pv, label):
    N = model['N']
    P_buy, P_ch, P_dis = sol['P_buy'], sol['P_ch'], sol['P_dis']
    E, P_pvu, P_cur = sol['E'], sol['P_pvu'], sol['P_cur']
    ok = True
    print(f"\n[验证:{label}]")

    # (1) 充放电互斥：min(P_ch, P_dis) <= τ（文档5.5）
    overlap = np.minimum(P_ch, P_dis)
    excl_ok = bool(np.all(overlap <= TAU))
    ok &= excl_ok
    print(f"  [{'PASS' if excl_ok else 'FAIL'}] 充放电互斥: max min(P_ch,P_dis)="
          f"{overlap.max():.3e} <= τ={TAU:.0e}")

    # (2) 功率平衡（文档4.2）
    bal = P_buy + P_pvu + P_dis - np.asarray(loads) - P_ch
    bal_ok = bool(np.max(np.abs(bal)) <= BAL_TOL)
    ok &= bal_ok
    print(f"  [{'PASS' if bal_ok else 'FAIL'}] 功率平衡: 最大残差 {np.max(np.abs(bal)):.3e} kW")

    # (3) 光伏平衡、弃光非负（文档4.3）
    pvbal = P_pvu + P_cur - np.asarray(pv)
    pv_ok = bool(np.max(np.abs(pvbal)) <= BAL_TOL
                 and np.min(P_pvu) >= -BAL_TOL and np.min(P_cur) >= -BAL_TOL)
    ok &= pv_ok
    print(f"  [{'PASS' if pv_ok else 'FAIL'}] 光伏平衡: 残差 {np.max(np.abs(pvbal)):.3e}, "
          f"P_pv_use∈[{P_pvu.min():.1f},{P_pvu.max():.1f}] kW, 弃光>=0")

    # (4) SOC范围与日初日末（文档4.4、第6节）
    soc_ok = bool(np.all(E >= E_MIN - 1e-6) and np.all(E <= E_MAX + 1e-6)
                  and abs(E[0] - E_INIT) <= BAL_TOL and abs(E[-1] - E_END) <= BAL_TOL)
    ok &= soc_ok
    print(f"  [{'PASS' if soc_ok else 'FAIL'}] SOC: [{E.min():.2f},{E.max():.2f}]"
          f"⊂[{E_MIN:.0f},{E_MAX:.0f}], E_1={E[0]:.4f}, E_145={E[-1]:.4f} kWh")

    # (5) 功率上界（文档4.5、4.6）
    pw_ok = bool(P_buy.max() <= P_BUY_MAX + 1e-6 and P_ch.max() <= P_CH_MAX + 1e-6
                 and P_dis.max() <= P_DIS_MAX + 1e-6)
    ok &= pw_ok
    print(f"  [{'PASS' if pw_ok else 'FAIL'}] 功率上界: 购电max={P_buy.max():.2f}, "
          f"充电max={P_ch.max():.2f}, 放电max={P_dis.max():.2f} kW")
    return ok


# ========== 5. 结果汇总（文档第6节） ==========
def print_summary(sol, model, prices, label):
    N = model['N']
    P_buy, P_ch, P_dis = sol['P_buy'], sol['P_ch'], sol['P_dis']
    E, P_pvu, P_cur = sol['E'], sol['P_pvu'], sol['P_cur']
    prices_arr = np.asarray(prices)

    total_buy = float(np.sum(P_buy) * DT)
    total_cost = float(np.dot(prices_arr, P_buy) * DT)
    total_cur = float(np.sum(P_cur) * DT)
    throughput = float(np.sum(P_ch + P_dis) * DT)
    total_pv_use = float(np.sum(P_pvu) * DT)

    print(f"\n[结果汇总:{label}]")
    print(f"  全天购电量   : {total_buy:.4f} kWh")
    print(f"  全天购电费   : {total_cost:.4f} 元")
    print(f"  弃光量       : {total_cur:.4f} kWh")
    print(f"  光伏利用量   : {total_pv_use:.4f} kWh")
    print(f"  储能吞吐量   : {throughput:.4f} kWh")
    print(f"  最大购电功率 : {P_buy.max():.2f} kW")
    print(f"  0:00 / 24:00 储电量: {E[0]:.4f} / {E[-1]:.4f} kWh")

    # C题表1指定时段（时段 p 覆盖 [p·10, (p+1)·10) 分钟，见文档第2节口径）
    table1 = [('10:00-10:10', 60), ('12:00-12:10', 72), ('14:00-14:10', 84),
              ('16:00-16:10', 96), ('18:00-18:10', 108), ('20:00-20:10', 120)]
    print("\n  表1指定时段购电量(kWh):")
    for lbl, p in table1:
        print(f"    {lbl}: {P_buy[p] * DT:.4f}")

    # C题表2分时段充放电量与储电量
    print("\n  表2分时段充放电量(kWh):")
    for s, e in [(0, 4), (4, 8), (8, 12), (12, 16), (16, 20), (20, 24)]:
        ch = float(np.sum(P_ch[s * 6:e * 6]) * DT)
        dis = float(np.sum(P_dis[s * 6:e * 6]) * DT)
        print(f"    {s}:00-{e}:00  充 {ch:.4f}  放 {dis:.4f}")

    # 小时级统计
    print("\n  小时级功率统计(kW):")
    for h in range(0, 24, 4):
        idx_start = h * 6
        idx_end = (h + 4) * 6
        if idx_end <= N:
            pb = P_buy[idx_start:idx_end]
            pc = P_ch[idx_start:idx_end]
            pd = P_dis[idx_start:idx_end]
            print(f"    {h}:00-{h+4}:00  购电[{pb.min():.1f},{pb.max():.1f}] 充电[{pc.min():.1f},{pc.max():.1f}] 放电[{pd.min():.1f},{pd.max():.1f}]")

    # SOC轨迹关键点
    print("\n  SOC轨迹关键点(kWh):")
    for t in [0, 24, 48, 72, 96, 120, 144]:
        print(f"    E_{t+1:>3d} = {E[t]:.2f}")


# ========== 6. 按附件5模板原结构写出 result1.xlsx ==========
def write_result(x, model, out_path, template_path):
    N = model['N']
    sol = extract(x, model)
    P_buy, P_ch, P_dis, E = sol['P_buy'], sol['P_ch'], sol['P_dis'], sol['E']

    wb = load_workbook(template_path)

    # Sheet1 计划购电量：按模板行序填入第 k 个时段的购电量(kWh)。
    # 口径（文档第2节）：附件1第 k 行数据（时刻 (k+1)·10 min）对应
    # 第 k 个时段；模板第 k 个数据行依次填入，保持模板标签不变。
    ws1 = wb['计划购电量']
    k = 0
    for r in range(2, ws1.max_row + 1):
        if ws1.cell(row=r, column=1).value in (None, ''):
            continue
        if k >= N:
            raise ValueError("模板'计划购电量'数据行超过144行，请检查附件5模板")
        ws1.cell(row=r, column=2).value = round(float(P_buy[k]) * DT, 4)
        k += 1
    if k != N:
        raise ValueError(f"模板'计划购电量'数据行仅{k}行，不足144行")

    # Sheet2 充放电量：按模板自带标签填充分时段充放电量与 0:00/24:00 储电量
    ws2 = wb['充放电量']
    chunks = {'0:00-4:00': (0, 4), '4:00-8:00': (4, 8), '8:00-12:00': (8, 12),
              '12:00-16:00': (12, 16), '16:00-20:00': (16, 20), '20:00-24:00': (20, 24)}
    for r in range(2, ws2.max_row + 1):
        lbl = ws2.cell(row=r, column=1).value
        if lbl in chunks:
            s, e = chunks[lbl]
            ws2.cell(row=r, column=2).value = round(float(np.sum(P_ch[s * 6:e * 6]) * DT), 4)
            ws2.cell(row=r, column=3).value = round(float(np.sum(P_dis[s * 6:e * 6]) * DT), 4)
        mark = ws2.cell(row=r, column=4).value
        if mark == '0:00':
            ws2.cell(row=r, column=5).value = round(float(E[0]), 4)
        elif mark == '24:00':
            ws2.cell(row=r, column=5).value = round(float(E[N]), 4)

    wb.save(out_path)
    print(f"\n[输出] 已按附件5模板写出 {out_path}")


# ========== 7. 按C题原文格式输出表1、表2 ==========
def write_table1(sol, model, prices, path):
    """
    C题原文表1（精确复刻）：
    |时间段|购电量|时间段|购电量|时间段|购电量|
    |10:00-10:10||12:00-12:10||14:00-14:10|
    |16:00-16:10||18:00-18:10||20:00-20:10|
    |全 天购电量||全 天购电费|||
    """
    N = model['N']
    P_buy = sol['P_buy']
    total_buy = float(np.sum(P_buy) * DT)
    total_cost = float(np.dot(np.asarray(prices), P_buy) * DT)
    slots = [('10:00-10:10', 60), ('12:00-12:10', 72), ('14:00-14:10', 84),
             ('16:00-16:10', 96), ('18:00-18:10', 108), ('20:00-20:10', 120)]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '表1'

    # 第1行：标题
    ws.merge_cells('A1:F1')
    ws['A1'] = '表1  微网在指定时间段的购电量及全天的购电量和购电费'
    ws['A1'].font = openpyxl.styles.Font(bold=True)

    # 第2行：表头
    for col, val in enumerate(['时间段', '购电量', '时间段', '购电量', '时间段', '购电量'], 1):
        ws.cell(row=2, column=col, value=val)
        ws.cell(row=2, column=col).font = openpyxl.styles.Font(bold=True)

    # 第3行：10:00-14:10
    for i, (lbl, idx) in enumerate(slots[:3]):
        ws.cell(row=3, column=2*i+1, value=lbl)
        ws.cell(row=3, column=2*i+2, value=round(float(P_buy[idx] * DT), 4))

    # 第4行：16:00-20:10
    for i, (lbl, idx) in enumerate(slots[3:]):
        ws.cell(row=4, column=2*i+1, value=lbl)
        ws.cell(row=4, column=2*i+2, value=round(float(P_buy[idx] * DT), 4))

    # 第5行：全天汇总
    ws.cell(row=5, column=1, value='全 天购电量')
    ws.cell(row=5, column=2, value=round(total_buy, 4))
    ws.cell(row=5, column=3, value='全 天购电费')
    ws.cell(row=5, column=4, value=round(total_cost, 4))

    wb.save(path)
    print(f"[输出] 表1 → {path}")


def write_table2(sol, model, path):
    """
    C题原文表2（精确复刻）：
    |时间段|充电量|放电量|时间段|充电量|放电量|
    |0:00-4:00|||4:00-8:00|||
    |8:00-12:00|||12:00-16:00|||
    |16:00-20:00|||20:00-24:00|||
    |0:00 储电量|||24:00 储电量|||
    """
    N = model['N']
    P_ch, P_dis, E = sol['P_ch'], sol['P_dis'], sol['E']
    pairs = [
        (('0:00-4:00', 0, 4), ('4:00-8:00', 4, 8)),
        (('8:00-12:00', 8, 12), ('12:00-16:00', 12, 16)),
        (('16:00-20:00', 16, 20), ('20:00-24:00', 20, 24)),
    ]

    def qty(pwr, s, e):
        return float(np.sum(pwr[s*6:e*6]) * DT)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '表2'

    # 第1行：标题
    ws.merge_cells('A1:F1')
    ws['A1'] = '表2  储能设备在指定时间段的充放电量及0:00和24:00的储电量'
    ws['A1'].font = openpyxl.styles.Font(bold=True)

    # 第2行：表头
    for col, val in enumerate(['时间段', '充电量', '放电量', '时间段', '充电量', '放电量'], 1):
        ws.cell(row=2, column=col, value=val)
        ws.cell(row=2, column=col).font = openpyxl.styles.Font(bold=True)

    # 第3-5行：时段数据
    for row_idx, ((lbl1, s1, e1), (lbl2, s2, e2)) in enumerate(pairs, start=3):
        ws.cell(row=row_idx, column=1, value=lbl1)
        ws.cell(row=row_idx, column=2, value=round(qty(P_ch, s1, e1), 4))
        ws.cell(row=row_idx, column=3, value=round(qty(P_dis, s1, e1), 4))
        ws.cell(row=row_idx, column=4, value=lbl2)
        ws.cell(row=row_idx, column=5, value=round(qty(P_ch, s2, e2), 4))
        ws.cell(row=row_idx, column=6, value=round(qty(P_dis, s2, e2), 4))

    # 第6行：储电量
    ws.cell(row=6, column=1, value='0:00 储电量')
    ws.cell(row=6, column=2, value=round(float(E[0]), 4))
    ws.cell(row=6, column=4, value='24:00 储电量')
    ws.cell(row=6, column=5, value=round(float(E[N]), 4))

    wb.save(path)
    print(f"[输出] 表2 → {path}")


# ========== 主程序（文档5.6 最终求解步骤） ==========
def main():
    print("=" * 60)
    print("问题一：微网日前购电策略求解（严格按建模思路文档实现）")
    print("=" * 60)

    # 步骤1：读取附件一的 c_t、L_t、P_pv_hat
    prices, loads, pv = read_data(DATA_FILE)
    N = len(prices)
    loads_arr = np.asarray(loads)
    pv_arr = np.asarray(pv)

    # 数据口径核对（文档第2节特征表）
    print(f"[数据] 全天负载电量={float(np.sum(loads_arr) * DT) / 1000:.3f} MWh, "
          f"全天预测光伏电量={float(np.sum(pv_arr) * DT) / 1000:.3f} MWh, "
          f"最大净负荷={float(np.max(loads_arr - pv_arr)):.1f} kW")

    # 步骤2：设置并确认 P_buy,max 与储能参数、初末SOC
    buy_max_str = '无上限' if np.isinf(P_BUY_MAX) else f'{P_BUY_MAX:.0f} kW'
    print(f"[参数] P_buy_max={buy_max_str}（题目未给变压器容量，理论上无上限）")
    print(f"[参数] 储能电量[{E_MIN:.0f},{E_MAX:.0f}] kWh, 初末SOC={E_INIT:.0f} kWh, "
          f"充/放功率<={P_CH_MAX:.0f} kW, η_ch=η_dis={ETA_CH}")

    # 步骤3：建立并求解第4节基准MILP（分支切割法）
    m_milp = build_constraints(N, loads, pv, with_u=True)
    milp_res, t_milp = solve_milp(m_milp, prices)
    milp_sol = extract(milp_res.x, m_milp)
    ok_milp = validate(milp_sol, m_milp, loads, pv, 'MILP基准')

    # 步骤4-5：删除 u_t 建立 LP 主模型，执行两阶段词典序优化
    m_lp = build_constraints(N, loads, pv, with_u=False)
    res1, res2, c_star, eps, t1, t2 = solve_lp_two_stage(m_lp, prices)
    lp_sol = extract(res2.x, m_lp)
    ok_lp = validate(lp_sol, m_lp, loads, pv, 'LP词典序终解')

    # 步骤7：比较LP与MILP的目标值和求解时间（文档5.5）
    # C_LP* = C_MILP* 的理论等式在浮点精度下允许微小相对偏差
    # 取相对容差 1e-5（约 0.36 元）作为数值验证阈值
    c_lp = float(cost_vector(m_lp, prices) @ res2.x)
    diff = abs(c_lp - float(milp_res.fun))
    rel_diff = diff / max(1.0, abs(c_lp))
    agree = rel_diff <= 1e-5
    print(f"\n[模型对比]")
    print(f"  C_LP*={c_lp:.4f} 元, C_MILP*={float(milp_res.fun):.4f} 元, |差|={diff:.4f} 元 (相对{rel_diff:.2e})")
    print(f"  求解耗时: MILP {t_milp:.2f}s vs LP {t1 + t2:.2f}s")
    print(f"  C_LP* = C_MILP* {'成立' if agree else '不成立，需检查'}（文档5.5）")

    # 步骤6：结果检验汇总与文件输出（采用词典序终解作为最终策略）
    print_summary(lp_sol, m_lp, prices, '最终策略(LP词典序终解)')
    write_result(res2.x, m_lp, OUT_FILE, TEMPLATE_FILE)
    write_table1(lp_sol, m_lp, prices, TABLE1_FILE)
    write_table2(lp_sol, m_lp, TABLE2_FILE)

    if not (ok_milp and ok_lp and agree):
        print("\n部分验证未通过!")
        sys.exit(1)
    print("\n全部验证通过!")


if __name__ == '__main__':
    main()

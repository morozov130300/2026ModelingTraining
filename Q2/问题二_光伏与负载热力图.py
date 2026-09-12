from pathlib import Path
import calendar
from datetime import date

import numpy as np
import openpyxl
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from matplotlib.colors import Normalize, LinearSegmentedColormap


# ---------- 路径（基于脚本位置定位，避免依赖运行目录 CWD）----------
BASE = Path(__file__).resolve().parent            # Q2/
XLSX = BASE.parent / "附件" / "附件2.xlsx"        # 附件/附件2.xlsx
OUTPUT = BASE / "annual_pv_2025.png"             # 图1（光伏环形）输出，固定落在 Q2/
OUTPUT2 = BASE / "annual_load_heatmap.png"       # 图2（负载 7×24）输出，同 Q2/

# ---------- 参数 ----------
YEAR = 2025
N_SLOTS = 144                                     # 附件2 每日期段列数（10 分钟 → 144）

# 季节定义（北半球气象季节，按月）
SEASON_OF_MONTH = {
    1: "冬", 2: "冬", 3: "春", 4: "春", 5: "春",
    6: "夏", 7: "夏", 8: "夏",
    9: "秋", 10: "秋", 11: "秋",
    12: "冬",
}
# 角度锚点：1 月 1 日在顶（theta=0°/N）、顺时针；每月严格 30°，季内按天均分
SEASON_ANGLE_DEG = {"冬": 15, "春": 105, "夏": 195, "秋": 285}   # 四季中心（四个对角）
SEASON_BOUND_DEG = [60, 150, 240, 330]                    # 季与季之间的分界线（粗）

# 色标：vmax 压到“白天有光照”区间，让夏季饱和到红端、冬季落在冷色端
VMAX_AUTO  = True        # True=按“冬季白天 P90”自动；False=用 VMAX_FIXED
VMAX_FIXED = 4000.0     # kW，VMAX_AUTO=False 时使用
DAYLIGHT   = slice(36, 114)   # 白天 06:00–18:50（10 分钟 slot 36..113），用于统计与定 vmax

# 2025 年法定节假日（国办发明电〔2024〕12号）：仅用于负载图旁标注；均值含全部天、不剔除
HOLIDAYS_2025 = ("元旦 1/1 · 春节 1/28–2/4 · 清明 4/4–4/6 · "
                 "五一 5/1–5/5 · 端午 5/31–6/2 · 国庆+中秋 10/1–10/8")

# 字体：宋体（Windows 有 SimSun）；基准字号放大到 15（默认 10 的 1.5 倍）
plt.rcParams["font.sans-serif"] = [
    "SimSun", "宋体", "Microsoft YaHei", "DejaVu Sans"
]
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["font.size"] = 15


# ---------- 读取附件2「光伏发电实际功率」宽表（365 x 144, kW）----------
# 与 solve_q2.py 的 read_actuals 一致口径：第1列日期，后 144 列为
# 时段终点功率（00:10 = 00:00–00:10，…，0:00+1 = 23:50–24:00）。
wb = openpyxl.load_workbook(XLSX, read_only=True)
pv_sheet = next(s for s in wb.sheetnames if "光伏" in s)
ws = wb[pv_sheet]

dates, rows = [], []
for row in ws.iter_rows(min_row=2, values_only=True):
    d0 = row[0]
    if d0 is None:
        continue
    vals = [float(v) for v in row[1:1 + N_SLOTS]]
    dates.append(d0.date())
    rows.append(vals)
wb.close()

values = np.asarray(rows, dtype=float)            # shape (365 天, 144 时段)
assert values.shape[1] == N_SLOTS, "附件2 时段列数与预期不符"

# 仅保留 YEAR 年数据（用标准库 date，无需 pandas）
start, end = date(YEAR, 1, 1), date(YEAR + 1, 1, 1)
keep = np.array([start <= d < end for d in dates])
values = values[keep]
dates = [d for d, k in zip(dates, keep) if k]
n_days, n_slots = values.shape

if n_days == 0:
    raise ValueError(f"没有找到 {YEAR} 年的光伏数据。")
if not np.isfinite(values).all():
    raise ValueError("光伏数据中存在缺失/无穷大，请先检查。")

months_of_day = np.array([d.month for d in dates], dtype=int)
month_days = {m: calendar.monthrange(YEAR, m)[1] for m in range(1, 13)}

print(f"数据源：{XLSX} [{pv_sheet}]")
print(f"日期数：{n_days}，每天时段数：{n_slots}")
print(f"光伏全年范围：[{values.min():.4f}, {values.max():.4f}] kW")

# 色标 vmax（数据驱动，让夏季饱和、冬季铺满下段 → 放大季节差异）
if VMAX_AUTO:
    winter_mask = np.isin(months_of_day, [12, 1, 2])
    winter_day = values[winter_mask][:, DAYLIGHT].ravel()
    winter_day = winter_day[winter_day > 0]
    VMAX = float(np.quantile(winter_day, 0.90))
else:
    VMAX = float(VMAX_FIXED)
print(f"色标范围：[0, {VMAX:.1f}] kW（{'自动=冬季白天 P90' if VMAX_AUTO else '固定'}）；0 kW=冷端（蓝）、强光=红端（蓝→红彩色过渡）")

# 每季白天典型功率（供报告引用“夏/冬典型峰谷”）
for sname in ["冬", "春", "夏", "秋"]:
    mset = [m for m, s in SEASON_OF_MONTH.items() if s == sname]
    day = values[np.isin(months_of_day, mset)][:, DAYLIGHT]
    print(f"  {sname}季白天功率：min {day.min():.1f} / 均值 {day.mean():.1f} / "
          f"P50 {np.quantile(day, 0.5):.1f} / P90 {np.quantile(day, 0.9):.1f} / "
          f"max {day.max():.1f} kW")


# ---------- 角度映射：每月严格 30°（“根据月份平均摆放”），季内按天均分 ----------
step = (2 * np.pi / 12) / np.array([month_days[m] for m in months_of_day])
theta_edges = np.concatenate([[0.0], np.cumsum(step)])
theta_edges = theta_edges / theta_edges[-1] * 2 * np.pi    # 归一到 [0, 2π]

# ---------- 径向网格：144 个 10 分钟时段 ----------
inner_radius, outer_radius = 1.5, 3.0
radius_edges = np.linspace(inner_radius, outer_radius, n_slots + 1)

# ---------- 数据 → 网格 + 冷→热彩色带（0 kW=冷端蓝，强光=热端红，全程有色相、无灰）----------
C = values.T                                                # (时段, 天)
Cm = C                                                     # 不 mask；0 落在冷端蓝，与色标图例一致

# 冷（深蓝）→ 蓝 → 青 → 黄绿 → 黄 → 橙 → 热（红/深红）；顶端红，无灰
cmap = LinearSegmentedColormap.from_list(
    "cold2hot_pv",
    ["#0b1e6f", "#1e88e5", "#26c6da", "#aeea00", "#ffd600", "#fb8c00", "#e53935", "#b71c1c"],
    N=256,
)
cmap.set_bad("#0b1e6f")   # 保险：任何残留 mask 也归到冷端蓝（非灰）
norm = Normalize(vmin=0.0, vmax=VMAX)


# ---------- 绘图（图1：光伏环形）----------
fig, ax = plt.subplots(
    figsize=(11, 10),
    subplot_kw={"projection": "polar"},
    layout="constrained",
)
ax.set_theta_zero_location("N")
ax.set_theta_direction(-1)
ax.set_ylim(0, outer_radius + 0.9)          # 外圈留出季节标签空间
ax.grid(False)
ax.spines["polar"].set_visible(False)

mesh = ax.pcolormesh(
    theta_edges, radius_edges, Cm,
    cmap=cmap, norm=norm,
    shading="flat", edgecolors="none", rasterized=True,
)


# ---------- 月分界（细）与季分界（粗）----------
for k in range(1, 12):                       # 每月 30° 细线
    a = np.deg2rad(30 * k)
    ax.plot([a, a], [inner_radius, outer_radius],
            color="white", linewidth=0.5, alpha=0.5)
for a in map(np.deg2rad, SEASON_BOUND_DEG):   # 季界粗线
    ax.plot([a, a], [inner_radius, outer_radius],
            color="white", linewidth=1.6, alpha=0.95)


# ---------- 季节标签（四个角）与月份小标签 ----------
lb_box = dict(facecolor="white", edgecolor="none", alpha=0.85, pad=2)
for name, deg in SEASON_ANGLE_DEG.items():    # 四季大标签（外）
    a = np.deg2rad(deg)
    ax.text(a, outer_radius + 0.42, name,
            ha="center", va="center",
            fontsize=27, fontweight="bold", color="#000000", bbox=lb_box)
for m in range(1, 13):                        # 月份小标签（内）
    a = np.deg2rad((m - 0.5) * 30)
    ax.text(a, outer_radius + 0.16, str(m),
            ha="center", va="center",
            fontsize=12, color="#000000", bbox=lb_box)


# ---------- 日内时间刻度 ----------
hours = np.array([0, 6, 12, 18, 24])
time_radii = inner_radius + hours / 24 * (outer_radius - inner_radius)
ax.set_yticks(time_radii)
ax.set_yticklabels([f"{h:02d}:00" for h in hours], fontsize=13.5, color="#000000")
ax.set_rlabel_position(80)

for label in ax.get_yticklabels():
    label.set_bbox(dict(facecolor="white", edgecolor="none", alpha=0.8, pad=1))

circle_angles = np.linspace(0, 2 * np.pi, 721)
for radius in time_radii:
    ax.plot(circle_angles, np.full_like(circle_angles, radius),
            color="white", linewidth=0.5, alpha=0.45)


# ---------- 标题、中心说明和色标 ----------
ax.text(0.5, 0.52, str(YEAR), transform=ax.transAxes,
        ha="center", va="center", fontsize=45, fontweight="bold")
ax.text(0.5, 0.44, "小区光伏\n内 → 外：00:00 → 24:00",
        transform=ax.transAxes, ha="center", va="center",
        fontsize=16.5, linespacing=1.7)
ax.set_title(f"{YEAR} 年小区光伏出力环形热力图",
             fontsize=27, pad=30)

colorbar = fig.colorbar(mesh, ax=ax, pad=0.10, shrink=0.72, aspect=28)
colorbar.set_label("光伏功率（kW）", fontsize=18)
colorbar.set_ticks(np.linspace(0, VMAX, 6))

ax.legend(handles=[Patch(facecolor="#0b1e6f", label="无光照 / 0 kW（冷端）")],
          loc="lower left", bbox_to_anchor=(-0.05, -0.05), frameon=False)

fig.savefig(OUTPUT, dpi=300, bbox_inches="tight", facecolor="white")    # 图1


# ============================================================
#  图2：小区负载「星期 × 整点」热力图（7 × 24，格 = 全年该星期该小时均值）
# ============================================================
wb_l = openpyxl.load_workbook(XLSX, read_only=True)
load_sheet = next(s for s in wb_l.sheetnames if "负载" in s)
ws_l = wb_l[load_sheet]
l_dates, l_rows = [], []
for row in ws_l.iter_rows(min_row=2, values_only=True):
    d0 = row[0]
    if d0 is None:
        continue
    l_dates.append(d0.date())
    l_rows.append([float(v) for v in row[1:1 + N_SLOTS]])
wb_l.close()

L_all = np.asarray(l_rows, dtype=float)          # (365 天, 144 时段)
assert L_all.shape[1] == N_SLOTS, "附件2 负载表时段列数与预期不符"
lkeep = np.array([start <= d < end for d in l_dates])
L = L_all[lkeep]
l_dates = [d for d, k in zip(l_dates, lkeep) if k]
if L.shape[0] == 0:
    raise ValueError(f"没有找到 {YEAR} 年的小区负载数据。")
if not np.isfinite(L).all():
    raise ValueError("负载数据中存在缺失/无穷大，请先检查。")

# 星期几（周一=0 … 周日=6）；每小时 6 个 10 分钟 → 该小时均值 → 全年该星期该小时算术平均
l_weekday = np.array([d.weekday() for d in l_dates])
L_hourly = L.reshape(L.shape[0], 24, 6).mean(axis=2)                 # (n_days, 24)
load_avg = np.empty((7, 24))
for wd in range(7):
    load_avg[wd] = L_hourly[l_weekday == wd].mean(axis=0)

LVMAX = float(np.quantile(L_hourly.ravel(), 0.90))
WD_NAMES = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]
print(f"\n[负载图] 数据源：{XLSX} [{load_sheet}]；7×24 格，格 = {YEAR} 全年该星期该小时均值")
print(f"[负载图] 色标范围：[0, {LVMAX:.1f}] kW（全年负载 P90）；低负载=冷端（蓝）、高峰=红端（蓝→红彩色过渡）")
for wd, wdn in enumerate(WD_NAMES):
    print(f"    {wdn}：日均 {load_avg[wd].mean():.1f} kW / 峰值 {load_avg[wd].max():.1f} kW")

cmap2 = LinearSegmentedColormap.from_list("cold2hot_load",
    ["#0b1e6f", "#1e88e5", "#26c6da", "#aeea00", "#ffd600", "#fb8c00", "#e53935", "#b71c1c"], N=256)
cmap2.set_bad("#0b1e6f")

fig2, ax2 = plt.subplots(figsize=(12, 7.5), layout="constrained")
im2 = ax2.imshow(load_avg, cmap=cmap2, norm=Normalize(0.0, LVMAX),
                 aspect="auto", interpolation="nearest")
ax2.set_xticks(np.arange(24) + 0.5)
ax2.set_xticklabels([f"{h:02d}" for h in range(24)], fontsize=13.5)
ax2.set_yticks(np.arange(7) + 0.5)
ax2.set_yticklabels(WD_NAMES, fontsize=27)
ax2.set_xlabel("一天中的时刻（整点；00 = 00:00–01:00 … 23 = 23:00–24:00）", fontsize=18)
ax2.set_title(f"{YEAR} 年小区负载星期 × 整点热力图",
              fontsize=27, pad=18)
cbar2 = fig2.colorbar(im2, ax=ax2, aspect=30)
cbar2.set_label("平均负载（kW）", fontsize=18)
cbar2.set_ticks(np.linspace(0, LVMAX, 6))
cbar2.ax.tick_params(labelsize=15)

# 节假日标注（图下方一行，说明这些天已并入对应星期均值、未单独剔除）
hol_line = f"2025 法定节假日（已并入对应星期均值，未剔除）：{HOLIDAYS_2025}"
fig2.text(0.5, -0.03, hol_line, transform=fig2.transFigure,
          ha="center", va="top", fontsize=15, color="#000000")

fig2.savefig(OUTPUT2, dpi=300, bbox_inches="tight", facecolor="white")   # 图2


# ---------- 收尾 ----------
plt.show()
print(f"\n图1（光伏）已保存到：{OUTPUT.resolve()}")
print(f"图2（负载）已保存到：{OUTPUT2.resolve()}")

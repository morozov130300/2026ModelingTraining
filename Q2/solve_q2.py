# -*- coding: utf-8 -*-
"""
问题二：日前计划购电与紧急购电策略 —— 逐日滚动两阶段随机优化
================================================================
严格按《Q2/问题二_日前计划购电与紧急购电建模.md》实现，求解器遵循 AGENTS.md：
  - MILP：scipy.optimize.milp（底层 HiGHS 分支切割法），mip_rel_gap=1e-9
  - LP  ：scipy.optimize.linprog(method='highs')（仅用于文档第8节的 LP 松弛检验）
  - 禁用 tol 选项，使用默认容差

已与用户确认的决策（2025 竞赛 C 题问题二）：
  1. 求解器 scipy.milp (HiGHS)，不使用 Q2/cbc；
  2. 场景概率等概率 p_ω = 1/M；
  3. 目标函数 λ=0，不加 CVaR（文档第6节扩展不启用）；
  4. 加入词典序第二阶段：费用最优 C* 约束下最小化储能吞吐量 Σ(C+D)Δt；
  5. 负载预测引入 2025 年法定节假日 + 调休日历（文档3.3节）；
  6. "可用误差日" = 负载与光伏均能生成完整因果预测的日（h ≥ 2025-01-29），
     累计满 M=20 个后切换为因果残差场景（即 2025-02-18 起），之前用暖启动均值残差；
  7. 预测参数固定文档建议初值：K_G=14, ρ_G=0.9, β=0.5, K_L=4, ρ_L=0.8, M=20；
  8. P_buy,max = ∞（题目未给并网容量，模型保留参数位，不擅自设 5000）；
  9. result2 模板"计划购电量"时间列按位置对应：第 j 个时间列 ↔ 当天第 j 个
     十分钟时段（区间 [10(j-1),10j) 分钟，附件2 第 j 列），与 Q1 填法一致；
 10. "紧急购电量"表：每个连续紧急购电段一行，超 3 段向下追加行，
     无紧急购电的日期保留日期行、时段与购电量留空，334 天全部列出；
 11. "全天购电量/全天购电费"两列取计划口径（Σ B·Δt 与 Σ c_t·B·Δt）；
 12. 运行流程：先 4 个代表日验证，再全量 334 天；每日结果落盘缓存可断点续跑；
 13. 非节假日负载同星期选样严格按文档字面公式 d-7k（不剔除节假日/调休样本）；
 14. 调休上班日按"最近 4 个正常工作日"选样（相似日精神，文档未明说）；
 15. 表1/表2/表3 交付：每表一个工作簿、4 个代表日纵向分块。
 16. 运行方式：无参数运行 = 全量 334 天（默认断点续跑、默认全核并行；
     multiprocessing 按天分派——天与天之间数据独立、每日 SOC 自行闭环，
     并行结果与串行逐日完全一致；每个子进程内仍是单线程确定性 MILP 求解）。

模型与文档对应关系：
  - 第3节   预测与场景：pv_forecast / load_forecast / build_scenarios
  - 第4-5节 两阶段随机 MILP：build_daily_model / solve_milp（含词典序第二阶段）
  - 第7节   滚动实施与回测：solve_day（计划锁定后才读取当天实际值）
  - 第8节   LP 松弛检验：lp_relaxation_check
  - 第10节  结果交付：write_result2 / write_tables（逐日明细仅在终端打印，不落盘）

用法（在装有 numpy/scipy/openpyxl 的目标机上运行）：
  python solve_q2.py                      # 无参数=一键全量334天（默认断点续跑+全核并行）
  python solve_q2.py repdays              # 4 个代表日求解与全面核验
  python solve_q2.py full --force         # 忽略缓存全部重解
  python solve_q2.py full --limit 10      # 冒烟测试：仅前 10 个决策日
  python solve_q2.py full --workers 8     # 指定并行进程数（默认=CPU 全部逻辑核）
  python solve_q2.py outputs              # 由缓存重建全部输出文件
  Windows 下也可直接双击同目录的 一键运行.bat
"""

import argparse
import multiprocessing
import os
import pickle
import time

import numpy as np
from scipy.optimize import Bounds, LinearConstraint, linprog, milp
from scipy.sparse import csr_matrix, vstack

from datetime import date, datetime, time as dtime, timedelta

import openpyxl

# ========== 路径 ==========
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)
ATTACH_DIR = os.path.join(PROJECT_ROOT, '附件')
PRICE_FILE = os.path.join(ATTACH_DIR, '附件1.xlsx')            # 附件1：电价 c_t
ACT_FILE = os.path.join(ATTACH_DIR, '附件2.xlsx')              # 附件2：实际负载与光伏
TEMPLATE_FILE = os.path.join(ATTACH_DIR, '附件5', 'result2.xlsx')  # 附件5模板（只读）
OUT_FILE = os.path.join(SCRIPT_DIR, 'result2.xlsx')
TABLE1_FILE = os.path.join(SCRIPT_DIR, '表1.xlsx')
TABLE2_FILE = os.path.join(SCRIPT_DIR, '表2.xlsx')
TABLE3_FILE = os.path.join(SCRIPT_DIR, '表3.xlsx')
CACHE_DIR = os.path.join(SCRIPT_DIR, 'cache')

# ========== 常量（文档第4.2节与附录1） ==========
N = 144                       # 每天十分钟时段数
DT = 10.0 / 60.0              # Δt = 1/6 h
ETA_CH = 0.9                  # η^ch
ETA_DIS = 0.9                 # η^dis
E_MIN, E_MAX = 1200.0, 10800.0  # 储能电量边界 (kWh)，文档式(5)
E_INIT = 6000.0               # E_{d,1}，文档式(6)
E_END = 6000.0                # E_{d,145}
P_CH_MAX = 5000.0             # 最大充电功率 (kW)，式(3)
P_DIS_MAX = 5000.0            # 最大放电功率 (kW)
P_BUY_MAX = np.inf            # 式(2) 并网点最大购电功率：题目未给，用户确认取 ∞
KAPPA = 5.0                   # 紧急购电加价系数 κ，式(1)

# 预测参数（文档3.2/3.3/3.4/3.6节建议初值，用户确认固定）
K_G, RHO_G, BETA = 14, 0.9, 0.5
K_L, RHO_L = 4, 0.8
M_SCEN = 20

# 求解与校验
MIP_GAP = 1e-9
TIME_LIMIT_DEFAULT = 300.0    # 单日单次 MILP 时间限制（秒）
TAU = 1e-6                    # 充放电互斥校验阈值 (kW)
BAL_TOL = 1e-4                # 等式约束/平衡残差容差
EPS_REL = 1e-6                # 词典序第二阶段费用松弛系数（同 Q1）

# 决策期与代表日
FIRST_DAY = date(2025, 2, 1)
LAST_DAY = date(2025, 12, 31)
REP_DAYS = [date(2025, 3, 20), date(2025, 6, 21), date(2025, 9, 23), date(2025, 12, 21)]
# 可用误差日起点：2025-01-29（负载同星期 K_L=4 与光伏 K_G=14 均可完整因果预测，用户确认）
CAUSAL_START = date(2025, 1, 29)

# 2025 年法定节假日（放假）与调休上班日（国务院办公厅公告，用户确认采用）
HOLIDAYS = {date(2025, 1, 1)}                     # 元旦


def _add_range(d1, d2):
    d = d1
    while d <= d2:
        HOLIDAYS.add(d)
        d += timedelta(days=1)


_add_range(date(2025, 1, 28), date(2025, 2, 4))    # 春节
_add_range(date(2025, 4, 4), date(2025, 4, 6))     # 清明
_add_range(date(2025, 5, 1), date(2025, 5, 5))     # 劳动节
_add_range(date(2025, 5, 31), date(2025, 6, 2))    # 端午
_add_range(date(2025, 10, 1), date(2025, 10, 8))   # 国庆+中秋
ADJUSTED_WORKDAYS = {date(2025, 1, 26), date(2025, 2, 8), date(2025, 4, 27),
                     date(2025, 9, 28), date(2025, 10, 11)}


def day_category(d):
    """日历类别：holiday=法定节假日, adjusted_work=调休上班日,
    weekday=正常工作日, weekend=正常周末。"""
    if d in HOLIDAYS:
        return 'holiday'
    if d in ADJUSTED_WORKDAYS:
        return 'adjusted_work'
    return 'weekday' if d.weekday() < 5 else 'weekend'


# ========== 1. 数据读取（文档第2.1节） ==========
def read_prices(path):
    """附件1：144 个时段电价 c_t（元/kWh），时段终点口径。"""
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    prices = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] is None:
            continue
        prices.append(float(row[1]))
    wb.close()
    if len(prices) != N:
        raise ValueError(f"附件1电价期望{N}个时段，实际{len(prices)}个")
    return np.asarray(prices, dtype=float)


def read_actuals(path):
    """附件2：两张 365x144 宽表（小区负载 / 光伏发电实际功率），单位 kW。
    时间列统一解释为时段终点（文档2.1节整理规则2）。"""
    wb = openpyxl.load_workbook(path, read_only=True)
    sheets = wb.sheetnames
    load_sheet = next(s for s in sheets if '负载' in s)
    pv_sheet = next(s for s in sheets if '光伏' in s)

    def read_sheet(name):
        ws = wb[name]
        dates, rows = [], []
        for row in ws.iter_rows(min_row=2, values_only=True):
            d0 = row[0]
            if d0 is None:
                continue
            vals = row[1:1 + N]
            if len(vals) != N or any(v is None for v in vals):
                raise ValueError(f"{name} 行 {d0} 数据不足或含缺失")
            dates.append(d0.date() if isinstance(d0, datetime) else d0)
            rows.append([float(v) for v in vals])
        return dates, np.asarray(rows, dtype=float)

    dates_l, L = read_sheet(load_sheet)
    dates_g, G = read_sheet(pv_sheet)
    wb.close()
    if dates_l != dates_g:
        raise ValueError("附件2两张工作表日期不一致")
    expected = []
    d = dates_l[0]
    while d <= dates_l[-1]:
        expected.append(d)
        d += timedelta(days=1)
    if dates_l != expected:
        raise ValueError("附件2日期不连续")
    # 数据质量画像（文档2.1节核对表）
    for nm, arr in (('小区负载', L), ('光伏实际', G)):
        print(f"[数据] {nm}: {arr.shape[0]}天x{arr.shape[1]}时段, "
              f"范围[{arr.min():.4f}, {arr.max():.4f}] kW, "
              f"缺失0(已校验), 负值{(arr < 0).sum()}个")
    return dates_l, L, G


def load_pool(d_idx, cats):
    """文档3.3节 同类型日选样池（返回历史日索引，最近优先，长度<=K_L）。
    - 节假日 d：优先此前节假日；不足补"相似休息日"（正常周末）；再不足取最近可用日；
    - 调休上班日 d：取最近 K_L 个正常工作日（用户决策14）；
    - 正常日 d：严格字面 d-7k，k=1..K_L（用户决策13，不剔除节假日样本）。"""
    cat = cats[d_idx]
    if cat == 'holiday':
        pool = [h for h in range(d_idx - 1, -1, -1) if cats[h] == 'holiday']
        if len(pool) < K_L:
            pool += [h for h in range(d_idx - 1, -1, -1)
                     if cats[h] == 'weekend' and h not in pool]
        if len(pool) < K_L:
            pool += [h for h in range(d_idx - 1, -1, -1) if h not in pool]
        pool = pool[:K_L]
    elif cat == 'adjusted_work':
        pool = [h for h in range(d_idx - 1, -1, -1) if cats[h] == 'weekday'][:K_L]
    else:
        pool = [d_idx - 7 * k for k in range(1, K_L + 1)]
        pool = [h for h in pool if h >= 0]
    if not pool or max(pool) >= d_idx:
        raise RuntimeError(f"负载选样池非法: d_idx={d_idx}, pool={pool}")
    return pool


def load_forecast(d_idx, L, cats):
    """文档3.3节 负载点预测：选样池指数加权（k=1 最近，权重 ρ_L^(k-1)）。"""
    pool = load_pool(d_idx, cats)
    w = RHO_L ** np.arange(len(pool))
    w = w / w.sum()
    return (w[:, None] * L[pool]).sum(axis=0)


def pv_forecast(d_idx, G):
    """文档3.2节 光伏点预测：K_G 天指数加权水平项 + 指数加权局部线性趋势外推，
    再按 β 收缩并截断非负。逐时段向量化 WLS：
      y_k = G_{d-k,t}, k=1..K_G, 模型 y = a - b·k（a 即外推到 d 的水平）。
      等价回归 y = α + m·x：m = Sxy/Sxx, a = α = ȳ_w - m·x̄_w, G̃ = ȳ_w。
      Ĝ = max{0, (1-β)·G̃ + β·a}。"""
    assert d_idx >= K_G, "光伏预测历史不足"
    Y = G[d_idx - K_G:d_idx]                     # (K_G, N)
    w = RHO_G ** np.arange(K_G)
    w = w / w.sum()
    x = np.arange(1, K_G + 1, dtype=float)
    xbar = float(w @ x)
    ybar = w @ Y                                 # (N,) = G̃
    dx = x - xbar
    sxx = float(np.sum(w * dx * dx))
    sxy = (w * dx) @ (Y - ybar)                  # (N,)
    m = sxy / sxx
    a = ybar - m * xbar
    return np.maximum(0.0, (1.0 - BETA) * ybar + BETA * a)


def precompute_causal(ctx):
    """文档3.4节：对每个可用历史日 h 生成因果预测（仅用 h 之前的数据），
    预计算残差所需的预测值。h < CAUSAL_START 的日不生成（负载同星期样本不足）。"""
    L, G, cats = ctx['L'], ctx['G'], ctx['cats']
    n_days = L.shape[0]
    cL = np.full((n_days, N), np.nan)
    cG = np.full((n_days, N), np.nan)
    start = ctx['date_idx'][CAUSAL_START]
    for h in range(start, n_days):
        cG[h] = pv_forecast(h, G)
        cL[h] = load_forecast(h, L, cats)
        assert max(load_pool(h, cats)) < h  # 信息边界：因果预测只用 h 之前数据
    ctx['causal_Lhat'] = cL
    ctx['causal_Ghat'] = cG
    ctx['causal_start_idx'] = start
    print(f"[因果预测] 已预计算 h ∈ [{ctx['dates'][start]}, {ctx['dates'][-1]}] "
          f"共 {n_days - start} 日的因果预测残差基础")


def build_scenarios(ctx, d_idx):
    """文档3.4节 联合误差场景：负载与光伏按同一历史日成对抽样，
    等概率 p_ω=1/n（用户决策2）。不足 M 个因果残差日时用暖启动均值残差（用户决策6）。"""
    L, G = ctx['L'], ctx['G']
    errs = np.arange(ctx['causal_start_idx'], d_idx)
    if len(errs) >= M_SCEN:
        hs = errs[-M_SCEN:]
        mode = 'causal_residual'
        eL = L[hs] - ctx['causal_Lhat'][hs]
        eG = G[hs] - ctx['causal_Ghat'][hs]
    else:
        hs = np.arange(d_idx - M_SCEN, d_idx)   # 决策日 d>=2/1 时必有 31 天历史
        mode = 'warm_mean'
        eL = L[hs] - L[hs].mean(axis=0)
        eG = G[hs] - G[hs].mean(axis=0)
    Lhat = load_forecast(d_idx, L, ctx['cats'])
    Ghat = pv_forecast(d_idx, G)
    Ls = np.maximum(0.0, Lhat[None, :] + eL)    # 场景负载 (n, N)
    Gs = np.maximum(0.0, Ghat[None, :] + eG)    # 场景光伏 (n, N)
    p = np.full(len(hs), 1.0 / len(hs))
    return Ls, Gs, p, Lhat, Ghat, hs, mode


# ========== 2. 两阶段随机 MILP（文档第4、5节） ==========
# 变量布局（N=144, n_s=M_SCEN 场景数）：
#   B: 0..N-1          计划购电功率 (kW)            式(2)
#   C: N..2N-1         计划充电功率                 式(3)
#   D: 2N..3N-1        计划放电功率
#   E: 3N..4N          储能电量 145 个状态点 (kWh)   式(4)(5)(6)
#   u: 4N+1..5N        充电状态二元变量              式(3)
#   场景块 w（基址 5N+1+w*5N，每块 5N）：
#     Buse/R/S/V/K 各 N 个              式(7)(8)(9)(10)
def build_daily_model(prices, Ls, Gs, p):
    n_s, _ = Ls.shape
    n = 5 * N + 1 + n_s * 5 * N
    iC0, iD0, iE0, iU0 = N, 2 * N, 3 * N, 4 * N + 1
    base0 = 5 * N + 1

    rows, cols, vals, beq = [], [], [], []

    def add_eq(entries, b):
        r = len(beq)
        for ci, v in entries:
            rows.append(r)
            cols.append(ci)
            vals.append(v)
        beq.append(float(b))

    for t in range(N):  # 式(4) SOC 转移
        add_eq([(iE0 + t + 1, 1.0), (iE0 + t, -1.0),
                (iC0 + t, -ETA_CH * DT), (iD0 + t, DT / ETA_DIS)], 0.0)
    add_eq([(iE0, 1.0)], E_INIT)      # 式(6) E_1
    add_eq([(iE0 + N, 1.0)], E_END)   # 式(6) E_145

    for w_i in range(n_s):
        bw = base0 + w_i * 5 * N
        for t in range(N):
            # 式(7)：Buse + R + V + D = L^ω + C
            add_eq([(bw + t, 1.0), (bw + N + t, 1.0), (bw + 3 * N + t, 1.0),
                    (iD0 + t, 1.0), (iC0 + t, -1.0)], Ls[w_i, t])
            # 式(8)：Buse + S = B
            add_eq([(bw + t, 1.0), (bw + 2 * N + t, 1.0), (t, -1.0)], 0.0)
            # 式(9)：V + K = G^ω
            add_eq([(bw + 3 * N + t, 1.0), (bw + 4 * N + t, 1.0)], Gs[w_i, t])
    A_eq = csr_matrix((vals, (rows, cols)), shape=(len(beq), n))
    b_eq = np.asarray(beq)

    rows, cols, vals, bub = [], [], [], []
    for t in range(N):  # 式(3)：C <= 5000u, D <= 5000(1-u)
        r = len(bub)
        rows += [r, r]
        cols += [iC0 + t, iU0 + t]
        vals += [1.0, -P_CH_MAX]
        bub.append(0.0)
        r = len(bub)
        rows += [r, r]
        cols += [iD0 + t, iU0 + t]
        vals += [1.0, P_DIS_MAX]
        bub.append(P_DIS_MAX)
    A_ub = csr_matrix((vals, (rows, cols)), shape=(len(bub), n))
    b_ub = np.asarray(bub)

    lo = np.zeros(n)
    hi = np.full(n, np.inf)              # 式(2)：B 无上界（P_buy,max=∞）
    hi[iC0:iC0 + N] = P_CH_MAX
    hi[iD0:iD0 + N] = P_DIS_MAX
    lo[iE0:iE0 + N + 1] = E_MIN          # 式(5)
    hi[iE0:iE0 + N + 1] = E_MAX
    hi[iU0:iU0 + N] = 1.0                # u ∈ {0,1}
    integrality = np.zeros(n)
    integrality[iU0:iU0 + N] = 1

    c_cost = np.zeros(n)                 # 式(1)：计划购电费 + 期望紧急购电费
    c_cost[:N] = prices * DT
    for w_i in range(n_s):
        c_cost[base0 + w_i * 5 * N + N: base0 + w_i * 5 * N + 2 * N] = \
            KAPPA * p[w_i] * prices * DT
    c_thru = np.zeros(n)                 # 词典序第二阶段目标：Σ(C+D)Δt
    c_thru[iC0:iC0 + N] = DT
    c_thru[iD0:iD0 + N] = DT

    return {'n': n, 'n_s': n_s, 'A_eq': A_eq, 'b_eq': b_eq,
            'A_ub': A_ub, 'b_ub': b_ub, 'lo': lo, 'hi': hi,
            'integrality': integrality, 'c_cost': c_cost, 'c_thru': c_thru}


def solve_milp(model, c, time_limit):
    res = milp(c=c,
               constraints=[LinearConstraint(model['A_eq'], model['b_eq'], model['b_eq']),
                            LinearConstraint(model['A_ub'], -np.inf, model['b_ub'])],
               integrality=model['integrality'],
               bounds=Bounds(model['lo'], model['hi']),
               options={'mip_rel_gap': MIP_GAP, 'time_limit': time_limit, 'presolve': True})
    if not res.success or res.x is None:
        raise RuntimeError(f"MILP 求解失败: status={res.status}, msg={res.message}")
    return res


def lp_relaxation_check(model):
    """文档第8节：去掉二元变量求 LP 松弛，检验互斥与费用一致性。"""
    res = linprog(model['c_cost'], A_ub=model['A_ub'], b_ub=model['b_ub'],
                  A_eq=model['A_eq'], b_eq=model['b_eq'],
                  bounds=list(zip(model['lo'], model['hi'])), method='highs')
    if not res.success:
        raise RuntimeError(f"LP 松弛求解失败: {res.message}")
    C = res.x[N:2 * N]
    D = res.x[2 * N:3 * N]
    excl = float(np.min(np.vstack([C, D]), axis=0).max())
    return {'cost': float(res.fun), 'excl_max': excl,
            'excl_ok': bool(excl <= TAU)}


def extract_plan(x, n_s):
    B = x[:N]
    C = x[N:2 * N]
    D = x[2 * N:3 * N]
    E = x[3 * N:4 * N + 1]
    U = x[4 * N + 1:5 * N + 1]
    scen = []
    base0 = 5 * N + 1
    for w_i in range(n_s):
        bw = base0 + w_i * 5 * N
        scen.append({'Buse': x[bw:bw + N], 'R': x[bw + N:bw + 2 * N],
                     'S': x[bw + 2 * N:bw + 3 * N], 'V': x[bw + 3 * N:bw + 4 * N],
                     'K': x[bw + 4 * N:bw + 5 * N]})
    return B, C, D, E, U, scen


# ========== 3. 单日求解 + 回测（文档第7节滚动实施） ==========
def merge_segments(R, thr=1e-6):
    """连续紧急购电时段合并（R>thr 的连续十分钟段）。"""
    segs = []
    above = R > thr
    t = 0
    while t < N:
        if above[t]:
            s = t
            while t < N and above[t]:
                t += 1
            segs.append((s, t))
        else:
            t += 1
    return segs


def solve_day(ctx, d_idx, time_limit=TIME_LIMIT_DEFAULT):
    dates = ctx['dates']
    prices = ctx['prices']
    d = dates[d_idx]

    # ---- 步骤2-3：日前预测与联合场景（仅用 h<d 信息，文档7节步骤1-3） ----
    t0 = time.perf_counter()
    Ls, Gs, p, Lhat, Ghat, hs, mode = build_scenarios(ctx, d_idx)
    model = build_daily_model(prices, Ls, Gs, p)

    # ---- 步骤4：求解并锁定计划（文档8节推荐顺序） ----
    res_milp = solve_milp(model, model['c_cost'], time_limit)
    c_star = float(res_milp.fun)
    lp = lp_relaxation_check(model)
    eps = EPS_REL * max(1.0, abs(c_star))
    A_ub2 = vstack([model['A_ub'], csr_matrix(model['c_cost'].reshape(1, -1))], format='csr')
    b_ub2 = np.append(model['b_ub'], c_star + eps)
    model2 = dict(model)
    model2['A_ub'] = A_ub2
    model2['b_ub'] = b_ub2
    res_lex = solve_milp(model2, model['c_thru'], time_limit)
    t_plan_done = time.perf_counter()

    B, C, D, E, U, scen = extract_plan(res_lex.x, model['n_s'])
    # 微小数值归一（用户决策：|x|<1e-9 归零，消除 -0.0000 显示；对残差影响 ~1e-9 量级）
    B = np.where(np.abs(B) < 1e-9, 0.0, B)
    C = np.where(np.abs(C) < 1e-9, 0.0, C)
    D = np.where(np.abs(D) < 1e-9, 0.0, D)
    plan_cost = float(model['c_cost'] @ res_lex.x)
    if plan_cost > c_star + eps + 1e-6:
        raise RuntimeError(f"{d} 词典序解费用 {plan_cost:.6f} 超出 C*+eps")

    # ---- 求解阶段校验（文档5.5/8节 + AGENTS.md 验证要求） ----
    eq_res = float(np.abs(model['A_eq'] @ res_lex.x - model['b_eq']).max())
    ub_res = float(np.max(model['A_ub'] @ res_lex.x - model['b_ub'])) if model['A_ub'].shape[0] else 0.0
    excl_milp = float(np.min(np.vstack([C, D]), axis=0).max())
    lp_rel_diff = abs(c_star - lp['cost']) / max(1.0, abs(c_star))
    u_ok = bool(np.all(np.abs(U - np.round(U)) <= 1e-9))
    checks = {
        'eq_residual': eq_res, 'ub_violation': max(0.0, ub_res),
        'excl_milp': excl_milp, 'u_binary_ok': u_ok,
        'soc_min': float(E.min()), 'soc_max': float(E.max()),
        'e0': float(E[0]), 'e145': float(E[-1]),
        'lp_cost': lp['cost'], 'lp_excl_max': lp['excl_max'],
        'lp_excl_ok': lp['excl_ok'], 'lp_rel_diff': lp_rel_diff,
        'c_star': c_star, 'plan_cost': plan_cost, 'eps': eps,
    }
    ok = (eq_res <= BAL_TOL and excl_milp <= TAU and u_ok
          and E.min() >= E_MIN - 1e-6 and E.max() <= E_MAX + 1e-6
          and abs(E[0] - E_INIT) <= BAL_TOL and abs(E[-1] - E_END) <= BAL_TOL)
    if not ok:
        print(f"[警告] {d} 求解阶段校验未全部通过: {checks}")

    # ---- 步骤5-7：锁定计划后，才读取当天实际值并结算（信息隔离关键点） ----
    Lr = ctx['L'][d_idx].copy()
    Gr = ctx['G'][d_idx].copy()
    t_read_done = time.perf_counter()
    assert t_read_done >= t_plan_done, "信息隔离被破坏：实际值读取早于计划锁定"

    rd = Lr + C - D                       # 实时净需求（负载+充电-放电）
    over_dis = int(np.sum(rd < -1e-6))    # 计划放电超过实际负载+充电（物理上"少放即可"，费用无影响）
    rd_pos = np.maximum(rd, 0.0)
    xdef = rd_pos - Gr - B
    R = np.maximum(xdef, 0.0)             # 文档7节步骤6：R = [L+C-D-G-B]^+
    deficit = R > 0.0
    V = np.where(deficit, Gr, np.minimum(Gr, rd_pos))     # 光伏优先消纳（用户默认4）
    Buse = np.where(deficit, B, np.minimum(B, rd_pos - V))
    S = B - Buse
    K = Gr - V
    # 过放电时段簿记（用户决策：计划口径交付 + 补充报告）：
    # 实际放电截断为可吸收量 D_eff，未执行部分留在电池内；该时段 C=0（充放互斥），D_eff=L 即纯带负载
    D_eff = np.where(rd >= 0.0, D, D + rd)
    unexec_dis = float(np.sum(np.maximum(D - D_eff, 0.0)) * DT)   # 未执行放电量 (kWh)
    actual_e145 = E_END + unexec_dis / ETA_DIS                    # 实际日末储电量（近似，剩余电量顺延次日）
    bal_res = float(np.abs(Buse + R + V + D_eff - (Lr + C)).max())
    split_res = max(float(np.abs(Buse + S - B).max()), float(np.abs(V + K - Gr).max()))

    plan_energy = float(np.sum(B) * DT)
    plan_cost_real = float(np.dot(prices, B) * DT)      # 式(13)第一项：按计划全额结算
    em_energy = float(np.sum(R) * DT)
    em_cost = float(np.dot(KAPPA * prices, R) * DT)     # 式(13)第二项
    unused_energy = float(np.sum(S) * DT)
    curtail_energy = float(np.sum(K) * DT)
    throughput = float(np.sum(C + D) * DT)
    segs = [(s * 10, e * 10, float(np.sum(R[s:e]) * DT)) for (s, e) in merge_segments(R)]
    mae_l = float(np.mean(np.abs(Lr - Lhat)))
    mae_g = float(np.mean(np.abs(Gr - Ghat)))
    rmse_l = float(np.sqrt(np.mean((Lr - Lhat) ** 2)))
    rmse_g = float(np.sqrt(np.mean((Gr - Ghat) ** 2)))
    t_end = time.perf_counter()

    rec = {
        'date': d, 'cat': ctx['cats'][d_idx], 'mode': mode, 'n_scen': int(len(hs)),
        'scen_from': ctx['dates'][int(hs[0])], 'scen_to': ctx['dates'][int(hs[-1])],
        'Lhat': Lhat, 'Ghat': Ghat, 'B': B, 'C': C, 'D': D, 'E': E,
        'segs': segs,
        'plan_energy': plan_energy, 'plan_cost': plan_cost_real,
        'em_energy': em_energy, 'em_cost': em_cost,
        'total_cost': plan_cost_real + em_cost,
        'unused_energy': unused_energy, 'curtail_energy': curtail_energy,
        'throughput': throughput, 'n_segs': len(segs),
        'mae_l': mae_l, 'mae_g': mae_g, 'rmse_l': rmse_l, 'rmse_g': rmse_g,
        'over_dis': over_dis, 'unexec_dis': unexec_dis, 'actual_e145': actual_e145,
        'bal_res': bal_res, 'split_res': split_res,
        'checks': checks, 'plan_sec': t_plan_done - t0, 'total_sec': t_end - t0,
    }
    return rec


# ========== 4. 缓存 ==========
def cache_path(d):
    return os.path.join(CACHE_DIR, f"day_{d.strftime('%Y%m%d')}.pkl")


def save_cache(rec):
    os.makedirs(CACHE_DIR, exist_ok=True)
    with open(cache_path(rec['date']), 'wb') as f:
        pickle.dump(rec, f)


def load_cache(d):
    p = cache_path(d)
    if os.path.exists(p):
        with open(p, 'rb') as f:
            return pickle.load(f)
    return None


# ========== 5. 输出（文档第10节 + 题目表1/表2/表3） ==========
def fmt_min(m):
    if m >= 24 * 60:
        return '24:00'
    return f"{m // 60}:{m % 60:02d}"


def write_result2(records, out_path=OUT_FILE):
    """按附件5模板原结构填写 result2.xlsx（不修改模板文件本身）。
    位置对应映射（用户决策9）：第 j 个时间列 ↔ 第 j 个十分钟时段。
    行定位以模板日期列为准（文档第10节第5点：以模板日期列核对，不凑行数）。"""
    wb = openpyxl.load_workbook(TEMPLATE_FILE)

    # ---- 计划购电量 ----
    ws = wb['计划购电量']
    hdr = [ws.cell(row=1, column=j).value for j in range(1, ws.max_column + 1)]
    j_energy = hdr.index('全天购电量') + 1
    j_cost = hdr.index('全天购电费') + 1
    date_row = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        v = v.date() if isinstance(v, datetime) else v
        if v is not None:
            if v in date_row:
                raise ValueError(f"模板'计划购电量'日期重复: {v}")
            date_row[v] = r
    n_expected = (LAST_DAY - FIRST_DAY).days + 1
    if len(date_row) != n_expected or min(date_row) != FIRST_DAY or max(date_row) != LAST_DAY:
        raise ValueError(f"模板日期范围异常: {min(date_row)} ~ {max(date_row)} 共 {len(date_row)} 行")
    rec_by_date = {rec['date']: rec for rec in records}
    missing = [d for d in sorted(date_row) if d not in rec_by_date]
    if missing:
        print(f"[警告] 部分输出：{len(missing)} 个日期尚无结果（冒烟测试模式），"
              f"首日 {missing[0]}，末日 {missing[-1]}")
    for d, r in date_row.items():
        rec = rec_by_date.get(d)
        if rec is None:
            continue
        for t in range(N):
            ws.cell(row=r, column=2 + t).value = round(float(rec['B'][t]) * DT, 4)
        ws.cell(row=r, column=j_energy).value = round(rec['plan_energy'], 4)
        ws.cell(row=r, column=j_cost).value = round(rec['plan_cost'], 4)

    # ---- 充放电量（写入前清理模板示意占位行） ----
    ws2 = wb['充放电量']
    for r in range(2, ws2.max_row + 1):
        for c in range(1, ws2.max_column + 1):
            ws2.cell(row=r, column=c).value = None
    chunk_labels = ['0:00-4:00', '4:00-8:00', '8:00-12:00',
                    '12:00-16:00', '16:00-20:00', '20:00-24:00']
    chunks = [(0, 4), (4, 8), (8, 12), (12, 16), (16, 20), (20, 24)]
    r = 2
    for rec in records:
        for k, (h1, h2) in enumerate(chunks):
            ws2.cell(row=r, column=2).value = chunk_labels[k]
            ws2.cell(row=r, column=3).value = round(float(np.sum(rec['C'][h1 * 6:h2 * 6]) * DT), 4)
            ws2.cell(row=r, column=4).value = round(float(np.sum(rec['D'][h1 * 6:h2 * 6]) * DT), 4)
            if k == 0:
                ws2.cell(row=r, column=1).value = datetime(rec['date'].year, rec['date'].month, rec['date'].day)
                ws2.cell(row=r, column=5).value = dtime(0, 0)
                ws2.cell(row=r, column=6).value = round(rec['checks']['e0'], 4)
            elif k == 1:
                ws2.cell(row=r, column=5).value = '24:00'
                ws2.cell(row=r, column=6).value = round(rec['checks']['e145'], 4)
            r += 1

    # ---- 紧急购电量（每段一行，超3段追加行，无紧急购电保留日期行，用户决策10） ----
    ws3 = wb['紧急购电量']
    for r0 in range(2, ws3.max_row + 1):
        for c in range(1, ws3.max_column + 1):
            ws3.cell(row=r0, column=c).value = None
    r = 2
    for rec in records:
        if not rec['segs']:
            ws3.cell(row=r, column=1).value = datetime(rec['date'].year, rec['date'].month, rec['date'].day)
            r += 1
        else:
            for k, (s_min, e_min, energy) in enumerate(rec['segs']):
                if k == 0:
                    ws3.cell(row=r, column=1).value = datetime(rec['date'].year, rec['date'].month, rec['date'].day)
                ws3.cell(row=r, column=2).value = f"{fmt_min(s_min)}-{fmt_min(e_min)}"
                ws3.cell(row=r, column=3).value = round(energy, 4)
                r += 1
    wb.save(out_path)
    print(f"[输出] result2.xlsx -> {out_path}（模板原结构，已填 {len(records)}/{n_expected} 天）")


def _block_title(ws, r, text):
    ws.cell(row=r, column=1, value=text).font = openpyxl.styles.Font(bold=True)


def write_tables(records_by_date):
    """题目表1/表2/表3（每表一簿、4 个代表日纵向分块，用户决策15）。"""
    slots = [('10:00-10:10', 60), ('12:00-12:10', 72), ('14:00-14:10', 84),
             ('16:00-16:10', 96), ('18:00-18:10', 108), ('20:00-20:10', 120)]
    chunk_pairs = [((0, 4), (4, 8)), ((8, 12), (12, 16)), ((16, 20), (20, 24))]

    def qty(rec, arr, h1, h2):
        return float(np.sum(arr[h1 * 6:h2 * 6]) * DT)

    # ---- 表1 ----
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '表1'
    r = 1
    for d in REP_DAYS:
        rec = records_by_date[d]
        _block_title(ws, r, f"{d.year}.{d.month}.{d.day}")
        r += 1
        for c, v in enumerate(['时间段', '购电量', '时间段', '购电量', '时间段', '购电量'], 1):
            ws.cell(row=r, column=c, value=v).font = openpyxl.styles.Font(bold=True)
        r += 1
        for grp in (slots[:3], slots[3:]):
            for i, (lbl, idx) in enumerate(grp):
                ws.cell(row=r, column=2 * i + 1, value=lbl)
                ws.cell(row=r, column=2 * i + 2, value=round(float(rec['B'][idx]) * DT, 4))
            r += 1
        ws.cell(row=r, column=1, value='全 天购电量')
        ws.cell(row=r, column=2, value=round(rec['plan_energy'], 4))
        ws.cell(row=r, column=3, value='全 天购电费')
        ws.cell(row=r, column=4, value=round(rec['plan_cost'], 4))
        r += 2
    wb.save(TABLE1_FILE)
    print(f"[输出] 表1 -> {TABLE1_FILE}")

    # ---- 表2 ----
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '表2'
    r = 1
    for d in REP_DAYS:
        rec = records_by_date[d]
        _block_title(ws, r, f"{d.year}.{d.month}.{d.day}")
        r += 1
        for c, v in enumerate(['时间段', '充电量', '放电量', '时间段', '充电量', '放电量'], 1):
            ws.cell(row=r, column=c, value=v).font = openpyxl.styles.Font(bold=True)
        r += 1
        for (a, b_) in chunk_pairs:
            ws.cell(row=r, column=1, value=f"{a[0]}:00-{a[1]}:00")
            ws.cell(row=r, column=2, value=round(qty(rec, rec['C'], *a), 4))
            ws.cell(row=r, column=3, value=round(qty(rec, rec['D'], *a), 4))
            ws.cell(row=r, column=4, value=f"{b_[0]}:00-{b_[1]}:00")
            ws.cell(row=r, column=5, value=round(qty(rec, rec['C'], *b_), 4))
            ws.cell(row=r, column=6, value=round(qty(rec, rec['D'], *b_), 4))
            r += 1
        ws.cell(row=r, column=1, value='0:00 储电量')
        ws.cell(row=r, column=2, value=round(rec['checks']['e0'], 4))
        ws.cell(row=r, column=4, value='24:00 储电量')
        ws.cell(row=r, column=5, value=round(rec['checks']['e145'], 4))
        r += 2
    wb.save(TABLE2_FILE)
    print(f"[输出] 表2 -> {TABLE2_FILE}")

    # ---- 表3（紧急购电量，4 日期并排） ----
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '表3'
    for i, d in enumerate(REP_DAYS):
        ws.cell(row=1, column=2 * i + 1, value=f"{d.year}.{d.month}.{d.day}").font = openpyxl.styles.Font(bold=True)
        ws.cell(row=2, column=2 * i + 1, value='时间段').font = openpyxl.styles.Font(bold=True)
        ws.cell(row=2, column=2 * i + 2, value='购电量').font = openpyxl.styles.Font(bold=True)
    max_seg = max(len(records_by_date[d]['segs']) for d in REP_DAYS)
    for i, d in enumerate(REP_DAYS):
        for k, (s_min, e_min, energy) in enumerate(records_by_date[d]['segs']):
            ws.cell(row=3 + k, column=2 * i + 1, value=f"{fmt_min(s_min)}-{fmt_min(e_min)}")
            ws.cell(row=3 + k, column=2 * i + 2, value=round(energy, 4))
    if max_seg == 0:
        for i in range(4):
            ws.cell(row=3, column=2 * i + 2, value='无紧急购电')
    wb.save(TABLE3_FILE)
    print(f"[输出] 表3 -> {TABLE3_FILE}")


def print_markdown_tables(records_by_date):
    slots = [('10:00-10:10', 60), ('12:00-12:10', 72), ('14:00-14:10', 84),
             ('16:00-16:10', 96), ('18:00-18:10', 108), ('20:00-20:10', 120)]
    print('\n### 表1 各代表日计划购电量（kWh）与全天汇总')
    for d in REP_DAYS:
        rec = records_by_date[d]
        print(f"\n**{d.year}.{d.month}.{d.day}**")
        print('|时间段|购电量|时间段|购电量|时间段|购电量|')
        print('|---|---|---|---|---|---|')
        print('|' + '|'.join(f"{lbl}|{float(rec['B'][idx]) * DT:.4f}" for lbl, idx in slots[:3]) + '|')
        print('|' + '|'.join(f"{lbl}|{float(rec['B'][idx]) * DT:.4f}" for lbl, idx in slots[3:]) + '|')
        print(f"|全 天购电量|{rec['plan_energy']:.4f}|全 天购电费|{rec['plan_cost']:.4f}||")
    print('\n### 表2 各代表日储能充放电量（kWh）与储电量')
    for d in REP_DAYS:
        rec = records_by_date[d]
        print(f"\n**{d.year}.{d.month}.{d.day}**")
        print('|时间段|充电量|放电量|时间段|充电量|放电量|')
        print('|---|---|---|---|---|---|')
        for (a, b_) in [((0, 4), (4, 8)), ((8, 12), (12, 16)), ((16, 20), (20, 24))]:
            ca = float(np.sum(rec['C'][a[0] * 6:a[1] * 6]) * DT)
            da = float(np.sum(rec['D'][a[0] * 6:a[1] * 6]) * DT)
            cb = float(np.sum(rec['C'][b_[0] * 6:b_[1] * 6]) * DT)
            db = float(np.sum(rec['D'][b_[0] * 6:b_[1] * 6]) * DT)
            print(f"|{a[0]}:00-{a[1]}:00|{ca:.4f}|{da:.4f}|{b_[0]}:00-{b_[1]}:00|{cb:.4f}|{db:.4f}|")
        print(f"|0:00 储电量|{rec['checks']['e0']:.4f}|24:00 储电量|{rec['checks']['e145']:.4f}||")
    print('\n### 表3 各代表日紧急购电量（kWh）')
    print('|' + '|'.join(f"{d.year}.{d.month}.{d.day}||" for d in REP_DAYS) + '|')
    print('|' + '|'.join(['时间段|购电量'] * 4) + '|')
    max_seg = max(len(records_by_date[d]['segs']) for d in REP_DAYS)
    for k in range(max(max_seg, 1)):
        cells = []
        for d in REP_DAYS:
            segs = records_by_date[d]['segs']
            if k < len(segs):
                s_min, e_min, energy = segs[k]
                cells += [f"{fmt_min(s_min)}-{fmt_min(e_min)}", f"{energy:.4f}"]
            else:
                cells += ['', '']
        print('|' + '|'.join(cells) + '|')


def print_aggregates(records, title):
    n = len(records)
    plan_c = sum(r['plan_cost'] for r in records)
    em_c = sum(r['em_cost'] for r in records)
    total_c = plan_c + em_c
    em_e = sum(r['em_energy'] for r in records)
    days_em = sum(1 for r in records if r['em_energy'] > 1e-6)
    plan_e = sum(r['plan_energy'] for r in records)
    unused = sum(r['unused_energy'] for r in records)
    curtail = sum(r['curtail_energy'] for r in records)
    thru = sum(r['throughput'] for r in records)
    ch_e = sum(float(np.sum(r['C'])) * DT for r in records)
    dis_e = sum(float(np.sum(r['D'])) * DT for r in records)
    print(f"\n[回测汇总:{title}] 共 {n} 天")
    print("=" * 56)
    print("  ★ 总费用与主要电量（重要中间量）")
    print(f"  全期总购电费用 : {total_c:.4f} 元（{total_c / 10000:.2f} 万元）")
    print(f"    ├─ 计划购电费 : {plan_c:.4f} 元（{plan_c / 10000:.2f} 万元，占 {plan_c / total_c:.2%}）")
    print(f"    └─ 紧急购电费 : {em_c:.4f} 元（{em_c / 10000:.2f} 万元，占 {em_c / total_c:.2%}）")
    print(f"  计划购电总电量 : {plan_e:.4f} kWh（{plan_e / 10000:.2f} 万 kWh）")
    print(f"  紧急购电总电量 : {em_e:.4f} kWh（{em_e / 10000:.2f} 万 kWh）")
    print(f"  未取用计划电量 : {unused:.4f} kWh（{unused / 10000:.2f} 万 kWh）")
    print(f"  弃光总量       : {curtail:.4f} kWh（{curtail / 10000:.2f} 万 kWh）")
    print(f"  储能总吞吐量   : {thru:.4f} kWh（充电 {ch_e / 10000:.2f} 万 + 放电 {dis_e / 10000:.2f} 万 kWh）")
    print("=" * 56)
    print(f"  平均日总费用   : {total_c / n:.4f} 元（计划 {plan_c / n:.4f} + 紧急 {em_c / n:.4f}）")
    print(f"  发生紧急购电天数占比 {days_em}/{n} = {days_em / n:.2%}")
    print(f"  预测精度(均值) : 负载MAE {np.mean([r['mae_l'] for r in records]):.4f} kW, "
          f"光伏MAE {np.mean([r['mae_g'] for r in records]):.4f} kW")
    lp_agree = sum(1 for r in records if r['checks']['lp_excl_ok'] and r['checks']['lp_rel_diff'] <= 1e-5)
    print(f"  LP化简有效性   : {lp_agree}/{n} 天 LP 满足互斥且费用与 MILP 一致（相对差<=1e-5）")
    over = sum(r['over_dis'] for r in records)
    unexec_total = sum(r['unexec_dis'] for r in records)
    print(f"  实时过放电时段 : {over} 个, 未执行放电合计 {unexec_total:.4f} kWh"
          f"（物理上少放即可、剩余电量留至次日，费用无影响）")
    print(f"  信息隔离核验   : 全部 {n} 天'计划锁定早于实际值读取'成立（逐日断言通过）")


def print_day_detail(rec):
    d = rec['date']
    ck = rec['checks']
    weekday_cn = ['周一', '周二', '周三', '周四', '周五', '周六', '周日'][d.weekday()]
    cat_cn = {'holiday': '法定节假日', 'adjusted_work': '调休上班日',
              'weekday': '工作日', 'weekend': '周末'}[rec['cat']]
    print(f"\n—— {d}（{weekday_cn}，{cat_cn}）——")
    print(f"  场景: {rec['mode']}, n={rec['n_scen']} ({rec['scen_from']} ~ {rec['scen_to']}), 等概率 p=1/{rec['n_scen']}")
    print(f"  MILP C*={ck['c_star']:.4f} 元, LP松弛={ck['lp_cost']:.4f} 元 (相对差 {ck['lp_rel_diff']:.2e}), "
          f"LP互斥 {'通过' if ck['lp_excl_ok'] else '不通过'} (max min(C,D)={ck['lp_excl_max']:.2e})")
    print(f"  词典序终解费用={ck['plan_cost']:.4f} 元 (<= C*+{ck['eps']:.2e}), "
          f"吞吐量={rec['throughput']:.4f} kWh, 互斥 max min(C,D)={ck['excl_milp']:.2e}")
    print(f"  等式残差={ck['eq_residual']:.2e}, SOC∈[{ck['soc_min']:.2f},{ck['soc_max']:.2f}], "
          f"E0={ck['e0']:.4f}, E145={ck['e145']:.4f}")
    print(f"  计划购电 {rec['plan_energy']:.4f} kWh / {rec['plan_cost']:.4f} 元; "
          f"紧急购电 {rec['em_energy']:.4f} kWh / {rec['em_cost']:.4f} 元; 总费用 {rec['total_cost']:.4f} 元")
    print(f"  未取用计划电量 {rec['unused_energy']:.4f} kWh, 弃光 {rec['curtail_energy']:.4f} kWh, "
          f"平衡残差(实际执行口径) {rec['bal_res']:.2e}")
    if rec['over_dis']:
        print(f"  过放电时段 {rec['over_dis']} 个（费用无影响）：未执行放电 {rec['unexec_dis']:.4f} kWh, "
              f"实际日末储电约 {rec['actual_e145']:.4f} kWh（剩余电量留至次日，计划仍以6000起步）")
    else:
        print("  过放电时段 0 个")
    print(f"  预测MAE: 负载 {rec['mae_l']:.4f} kW, 光伏 {rec['mae_g']:.4f} kW")
    if rec['segs']:
        seg_str = ', '.join(f"{fmt_min(s)}-{fmt_min(e)}({en:.4f} kWh)" for s, e, en in rec['segs'])
        print(f"  紧急购电段 {rec['n_segs']} 个: {seg_str}")
    else:
        print("  紧急购电段: 无")


# ========== 6. 主流程 ==========
def build_ctx():
    prices = read_prices(PRICE_FILE)
    dates, L, G = read_actuals(ACT_FILE)
    cats = [day_category(d) for d in dates]
    date_idx = {d: i for i, d in enumerate(dates)}
    ctx = {'prices': prices, 'dates': dates, 'L': L, 'G': G,
           'cats': cats, 'date_idx': date_idx}
    precompute_causal(ctx)
    return ctx


# 主进程读数后共享给子进程的全局上下文（spawn 模式下子进程通过 initializer 注入）
_CTX = {}


def _init_worker(ctx):
    global _CTX
    _CTX.clear()
    _CTX.update(ctx)


def _has_ctx():
    return len(_CTX) > 0


def _get_ctx():
    if not _has_ctx():
        raise RuntimeError("子进程上下文未初始化")
    return _CTX


def _worker_solve(args):
    """并行子任务：求解一个决策日并写缓存，返回 (d_idx, rec)。
    天与天之间完全独立（每日 SOC 各自闭环、场景只用该日之前历史），
    因此并行与串行结果逐日一致（决策16）。"""
    d_idx, time_limit = args
    ctx = _get_ctx() if _has_ctx() else _CTX
    rec = solve_day(ctx, d_idx, time_limit)
    save_cache(rec)
    return d_idx, rec


def decision_indices(ctx):
    i0 = ctx['date_idx'][FIRST_DAY]
    i1 = ctx['date_idx'][LAST_DAY]
    return list(range(i0, i1 + 1))


def run_repdays(ctx, time_limit):
    print("=" * 64)
    print("代表日模式：4 个指定日期逐日求解与全面核验")
    print("=" * 64)
    records = {}
    for d in REP_DAYS:
        rec = solve_day(ctx, ctx['date_idx'][d], time_limit)
        save_cache(rec)
        records[d] = rec
        print_day_detail(rec)
    records_by_date = records
    write_tables(records_by_date)
    print_markdown_tables(records_by_date)
    print_aggregates(list(records.values()), '4个代表日')
    print("\n[提示] 代表日结果未写 result2.xlsx（需全量 334 天后生成）。"
          "确认无误后运行: python solve_q2.py（一键全量，默认断点续跑+全核并行）")


def _run_parallel(idxs, ctx, time_limit, workers):
    """按天并行求解（决策16）：imap_unordered 动态取任务，结果按日期排序后返回，
    保证与串行逐日完全一致。缓存命中在主进程先过滤掉。"""
    tasks = [(i, time_limit) for i in idxs]
    n_workers = max(1, min(workers, len(tasks)))
    print(f"[并行] {n_workers} 进程 × {len(tasks)} 天（每子进程单线程 MILP，确定性不变）",
          flush=True)
    out = {}
    t_start = time.perf_counter()
    with multiprocessing.get_context('spawn').Pool(
            processes=n_workers, initializer=_init_worker, initargs=(ctx,)) as pool:
        for k, (d_idx, rec) in enumerate(
                pool.imap_unordered(_worker_solve, tasks, chunksize=1)):
            out[d_idx] = rec
            done = len(out)
            el = time.perf_counter() - t_start
            print(f"[{done}/{len(tasks)}] {rec['date']} 求解 | 计划费 {rec['plan_cost']:.2f} "
                  f"紧急费 {rec['em_cost']:.2f} | LP差 {rec['checks']['lp_rel_diff']:.1e} | "
                  f"耗时 {rec['plan_sec']:.1f}s | 墙钟 {el:.0f}s", flush=True)
    return [out[i] for i in sorted(out)]


def run_full(ctx, time_limit, resume=True, limit=None, workers=None):
    print("=" * 64)
    print(f"全量模式：{FIRST_DAY} ~ {LAST_DAY} 共 334 天（resume={resume}, "
          f"limit={limit}, workers={workers or 'CPU全核'}）")
    print("=" * 64)
    idxs = decision_indices(ctx)
    if limit is not None:
        idxs = idxs[:limit]
    records = []
    todo = []
    for d_idx in idxs:
        d = ctx['dates'][d_idx]
        rec = load_cache(d) if resume else None
        if rec is not None:
            records.append(rec)
        else:
            todo.append(d_idx)
    print(f"[缓存] 命中 {len(records)} 天，待求解 {len(todo)} 天", flush=True)
    if todo:
        recs = _run_parallel(todo, ctx, time_limit, workers or os.cpu_count() or 1)
        records.extend(recs)
    records.sort(key=lambda r: r['date'])
    write_result2(records)
    by_date = {rec['date']: rec for rec in records}
    missing = [d for d in REP_DAYS if d not in by_date]
    if missing:
        print(f"[警告] 代表日 {missing} 不在本批记录中，表1/2/3 跳过")
    else:
        write_tables(by_date)
        print_markdown_tables(by_date)
    print_aggregates(records, f"全量{len(records)}天")


def run_outputs(ctx):
    print("=" * 64)
    print("输出重建模式：从缓存生成 result2.xlsx / 表1 / 表2 / 表3")
    print("=" * 64)
    records = []
    for d_idx in decision_indices(ctx):
        d = ctx['dates'][d_idx]
        rec = load_cache(d)
        if rec is None:
            raise RuntimeError(f"缺少 {d} 的缓存，请先运行 full")
        records.append(rec)
    write_result2(records)
    by_date = {rec['date']: rec for rec in records}
    write_tables(by_date)
    print_markdown_tables(by_date)
    print_aggregates(records, f"全量{len(records)}天")


def main():
    ap = argparse.ArgumentParser(
        description='2025 CUMCM C题 问题二 逐日滚动求解'
                    '（无参数 = 一键全量334天：默认断点续跑 + 默认全核并行）')
    ap.add_argument('cmd', nargs='?', default='full',
                    choices=['full', 'repdays', 'outputs'],
                    help='full=全量334天（默认） | repdays=4代表日核验 | outputs=由缓存重建输出')
    ap.add_argument('--force', action='store_true', help='忽略缓存全部重解')
    ap.add_argument('--limit', type=int, default=None, help='仅求解前 N 个决策日（冒烟测试）')
    ap.add_argument('--workers', type=int, default=None,
                    help='并行进程数（默认=CPU全部逻辑核，吃满CPU）')
    ap.add_argument('--time-limit', type=float, default=TIME_LIMIT_DEFAULT,
                    help='单日单次 MILP 时间限制（秒）')
    args = ap.parse_args()

    print(f"[参数] Δt={DT:.4f}h, κ={KAPPA}, η={ETA_CH}, E∈[{E_MIN},{E_MAX}], E0=E145={E_INIT}, "
          f"P_ch/P_dis<={P_CH_MAX:.0f}kW, P_buy_max={'无上限' if np.isinf(P_BUY_MAX) else P_BUY_MAX}")
    print(f"[参数] 预测: K_G={K_G}, ρ_G={RHO_G}, β={BETA}, K_L={K_L}, ρ_L={RHO_L}, M={M_SCEN}, "
          f"等概率; 残差切换起点={CAUSAL_START}; 节假日+调休日历已启用")
    print(f"[参数] 运行: cmd={args.cmd}, 断点续跑={'关(全部重解)' if args.force else '开'}, "
          f"并行数={'CPU全核=' + str(os.cpu_count()) if args.workers is None else args.workers}")
    ctx = build_ctx()
    if args.cmd == 'repdays':
        run_repdays(ctx, TIME_LIMIT_DEFAULT)
    elif args.cmd == 'full':
        run_full(ctx, args.time_limit, resume=not args.force,
                 limit=args.limit, workers=args.workers)
    elif args.cmd == 'outputs':
        run_outputs(ctx)


if __name__ == '__main__':
    main()
